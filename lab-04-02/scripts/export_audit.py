import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(__file__).resolve().parent.parent
audit = json.loads((root / "output/audit.checked.json").read_text())
if "离线" not in audit["provenance"]:
    raise ValueError("该导出器仅生成明确标识的离线教学夹具报告")
wb = Workbook()
ws = wb.active
ws.title = "Audit"
headers = ["合同编号", "条款", "合同原文", "规则编号", "规则原文", "结果",
           "原因", "复核状态", "ChunkId", "来源URL", "执行性质"]
ws.append(headers)
for row in audit["findings"]:
    values = [audit["contractId"], row["clauseId"], row["contractQuote"],
              row["ruleId"], row["ruleQuote"], row["result"], row["reason"],
              row["reviewStatus"], row["chunkId"], row["sourceUrl"], "离线教学夹具"]
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.data_type = "s"
table = Table(displayName="AuditRows", ref=f"A1:K{ws.max_row}")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(table)
ws.freeze_panes = "A2"
for cells in ws:
    for cell in cells:
        cell.font = Font(name="Arial", size=11, bold=cell.row == 1)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if cell.row == 1:
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
for column, width in {"A": 20, "B": 12, "C": 55, "D": 20, "E": 60, "F": 15,
                      "G": 50, "H": 20, "I": 35, "J": 45, "K": 22}.items():
    ws.column_dimensions[column].width = width
for index in range(2, ws.max_row + 1):
    ws.row_dimensions[index].height = 65
notes = wb.create_sheet("Readme")
notes.append(["性质", audit["provenance"]])
notes.append(["来源", "data/contract.md + data/policy.raw.md；全部虚构教学内容"])
notes.append(["引用URL", "example.invalid 为占位URL，不能作为真实可访问知识引用"])
notes.append(["公式", "本报告仅含文本记录，无计算公式"])
for row in notes:
    for cell in row:
        cell.font = Font(name="Arial", size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
notes.column_dimensions["A"].width = 18
notes.column_dimensions["B"].width = 100
target = root / "output/contract-audit.example.xlsx"
wb.save(target)
check = load_workbook(target)
assert check["Audit"].max_row == len(audit["findings"]) + 1
assert "AuditRows" in check["Audit"].tables
errors = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}
for sheet in check:
    for row in sheet:
        for cell in row:
            assert cell.data_type != "f" and cell.value not in errors
print("已保存并验证离线Excel审计报告：", target)
